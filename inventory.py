import argparse, csv, datetime as dt, json, logging, random, sqlite3, time
from pathlib import Path
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

SCOPE='https://www.googleapis.com/auth/drive.readonly'
FIELDS='nextPageToken,incompleteSearch,files(id,name,mimeType,size,createdTime,modifiedTime,parents,webViewLink,description,trashed,md5Checksum,owners(displayName,emailAddress),lastModifyingUser(displayName,emailAddress))'

def call(req, attempts=10):
    for n in range(attempts):
        try: return req.execute()
        except HttpError as e:
            reasons=set()
            if e.resp.status == 403:
                try:
                    details=json.loads(e.content).get('error',{}).get('errors',[])
                    reasons={item.get('reason') for item in details if isinstance(item,dict)}
                except (ValueError,TypeError,AttributeError):
                    pass
            retryable=e.resp.status in (429,500,502,503,504) or (
                e.resp.status == 403 and bool(reasons & {'userRateLimitExceeded','rateLimitExceeded'}))
            if not retryable or n==attempts-1: raise
            delay=min(2**n+random.random(),64)
            logging.warning('API HTTP %s (%s); nova tentativa %d/%d em %.1fs',
                            e.resp.status,','.join(sorted(r for r in reasons if r)) or 'transitório',n+2,attempts,delay)
            time.sleep(delay)

def database(path):
    db=sqlite3.connect(path)
    db.execute('PRAGMA journal_mode=WAL')
    db.executescript('''CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, started TEXT, finished TEXT, status TEXT, discovered INTEGER DEFAULT 0, scanned INTEGER DEFAULT 0, failed INTEGER DEFAULT 0);
    CREATE TABLE IF NOT EXISTS drives(id TEXT PRIMARY KEY,name TEXT, last_seen_run INTEGER);
    CREATE TABLE IF NOT EXISTS files(drive_id TEXT,id TEXT,name TEXT,mime TEXT,size INTEGER,created TEXT,modified TEXT,parent_ids TEXT,url TEXT,md5 TEXT,last_seen_run INTEGER,PRIMARY KEY(drive_id,id));
    CREATE INDEX IF NOT EXISTS idx_files_drive ON files(drive_id);
    CREATE TABLE IF NOT EXISTS drive_scans(run_id INTEGER,drive_id TEXT,status TEXT,error TEXT,PRIMARY KEY(run_id,drive_id));''')
    return db

def csv_out(db, output):
    output.mkdir(parents=True,exist_ok=True)
    query='''SELECT d.name,f.drive_id,f.id,f.name,f.mime,f.size,f.created,f.modified,f.parent_ids,f.url,f.md5,f.last_seen_run FROM files f JOIN drives d ON d.id=f.drive_id ORDER BY d.name,f.name'''
    with (output/'arquivos.csv').open('w',encoding='utf-8-sig',newline='') as fh:
        w=csv.writer(fh);w.writerow(['drive','drive_id','file_id','nome','mime','bytes','criado','modificado','pais_ids','url','md5','ultima_execucao']);w.writerows(db.execute(query))
    summary=list(db.execute('''SELECT d.name,d.id,COUNT(f.id),COALESCE(SUM(CASE WHEN f.mime!='application/vnd.google-apps.folder' THEN 1 ELSE 0 END),0),COALESCE(SUM(CASE WHEN f.mime='application/vnd.google-apps.folder' THEN 1 ELSE 0 END),0),COALESCE(SUM(f.size),0) FROM drives d LEFT JOIN files f ON f.drive_id=d.id GROUP BY d.id ORDER BY d.name'''))
    with (output/'resumo.csv').open('w',encoding='utf-8-sig',newline='') as fh:
        w=csv.writer(fh);w.writerow(['drive','drive_id','itens','arquivos','pastas','bytes','GiB']);w.writerows([list(r)+[round(r[5]/1024**3,3)] for r in summary])
    return summary

def xlsx_out(summary, output):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); sh=wb.active; sh.title='Resumo'
    rows=[['Drive','ID','Itens','Arquivos','Pastas','Bytes','GiB']]+[list(r)+[round(r[5]/1024**3,3)] for r in summary]
    for row in rows: sh.append(row)
    for cell in sh[1]:
        cell.fill=PatternFill(fill_type='solid',fgColor='12304A')
        cell.font=Font(bold=True,color='FFFFFF')
    sh.row_dimensions[1].height=27
    for col,width in [('A',35),('B',38),('C',14),('D',14),('E',14),('F',22),('G',15)]: sh.column_dimensions[col].width=width
    sh.freeze_panes='A2'
    wb.save(output/'resumo.xlsx')

def main():
    p=argparse.ArgumentParser(description='Inventário dos Drives Compartilhados do Google Workspace')
    p.add_argument('--key',required=True);p.add_argument('--admin',required=True);p.add_argument('--db',default='inventory.sqlite3');p.add_argument('--out',default='relatorios');p.add_argument('--no-xlsx',action='store_true')
    args=p.parse_args();logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s')
    creds=service_account.Credentials.from_service_account_file(args.key,scopes=[SCOPE]).with_subject(args.admin)
    api=build('drive','v3',credentials=creds,cache_discovery=False)
    db=database(args.db); now=lambda:dt.datetime.now(dt.timezone.utc).isoformat()
    cur=db.execute('INSERT INTO runs(started,status) VALUES (?,?)',(now(),'running'));run=cur.lastrowid;db.commit()
    drives=[];token=None
    try:
        while True:
            res=call(api.drives().list(pageSize=100,pageToken=token,fields='nextPageToken,drives(id,name)',useDomainAdminAccess=True))
            drives+=res.get('drives',[]);token=res.get('nextPageToken')
            if not token:break
    except Exception:
        db.execute('UPDATE runs SET status=?,finished=? WHERE id=?',('discovery_failed',now(),run));db.commit();raise
    db.execute('UPDATE runs SET discovered=? WHERE id=?',(len(drives),run));db.commit()
    failed=0
    for index,d in enumerate(drives,1):
        logging.info('Drive %s/%s: %s',index,len(drives),d['name'])
        db.execute('INSERT INTO drives(id,name,last_seen_run) VALUES(?,?,?) ON CONFLICT(id) DO UPDATE SET name=excluded.name,last_seen_run=excluded.last_seen_run',(d['id'],d['name'],run));db.commit()
        token=None;count=0
        try:
            while True:
                res=call(api.files().list(corpora='drive',driveId=d['id'],includeItemsFromAllDrives=True,supportsAllDrives=True,spaces='drive',q='trashed = false',pageSize=1000,pageToken=token,fields=FIELDS))
                if res.get('incompleteSearch'):raise RuntimeError('Google retornou incompleteSearch=true')
                batch=[(d['id'],f['id'],f.get('name'),f.get('mimeType'),int(f['size']) if f.get('size') else None,f.get('createdTime'),f.get('modifiedTime'),','.join(f.get('parents',[])),f.get('webViewLink'),f.get('md5Checksum'),run) for f in res.get('files',[])]
                db.executemany('''INSERT INTO files(drive_id,id,name,mime,size,created,modified,parent_ids,url,md5,last_seen_run) VALUES(?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(drive_id,id) DO UPDATE SET name=excluded.name,mime=excluded.mime,size=excluded.size,created=excluded.created,modified=excluded.modified,parent_ids=excluded.parent_ids,url=excluded.url,md5=excluded.md5,last_seen_run=excluded.last_seen_run''',batch)
                count+=len(batch);db.commit();token=res.get('nextPageToken')
                if not token:break
            # Remove arquivos ausentes SOMENTE após varredura integral do Drive.
            db.execute('DELETE FROM files WHERE drive_id=? AND last_seen_run!=?',(d['id'],run))
            db.execute('INSERT INTO drive_scans VALUES (?,?,?,?)',(run,d['id'],'ok',None));db.commit()
            logging.info('  %s itens',count)
        except Exception as e:
            failed+=1;db.execute('INSERT INTO drive_scans VALUES (?,?,?,?)',(run,d['id'],'failed',str(e)));db.commit();logging.error('  Falhou: %s',e)
    db.execute('UPDATE runs SET finished=?,status=?,scanned=?,failed=? WHERE id=?',(now(),'partial' if failed else 'complete',len(drives)-failed,failed,run));db.commit()
    out=Path(args.out);summary=csv_out(db,out)
    if not args.no_xlsx:
        try:xlsx_out(summary,out)
        except Exception as e:logging.warning('Excel indisponível; CSV gerado: %s',e)
    logging.info('Concluído. Drives descobertos: %d; falhas: %d; relatórios: %s',len(drives),failed,out.resolve())
    if failed:raise SystemExit(2)

if __name__=='__main__':main()
